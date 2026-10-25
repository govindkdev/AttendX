import csv
import io
from datetime import datetime, date, time, timezone
from typing import Optional, Any
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from app.db.session import get_db
from app.models.entities import (
    User,
    UserRole,
    StudentProfile,
    TeacherProfile,
    AttendanceRecord,
    AttendanceStatus,
    VerificationMethod,
    LectureSession,
    Subject,
    Course,
    Branch
)
from app.permissions.deps import get_current_user

router = APIRouter(prefix="/reports", tags=["Reports"])

def sanitize_formula(val: Any) -> str:
    """Prevents CSV/Excel formula injection (CWE-1236)."""
    s = str(val or "")
    if s.startswith(("=", "+", "-", "@")):
        return "'" + s
    return s

async def fetch_filtered_records(
    db: AsyncSession,
    current_user: User,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    subject_id: Optional[int] = None,
    status_filter: Optional[AttendanceStatus] = None,
    course_id: Optional[int] = None,
    branch_id: Optional[int] = None
):
    stmt = select(AttendanceRecord).options(
        selectinload(AttendanceRecord.lecture_session).selectinload(LectureSession.subject),
        selectinload(AttendanceRecord.student).selectinload(StudentProfile.user),
        selectinload(AttendanceRecord.student).selectinload(StudentProfile.course),
        selectinload(AttendanceRecord.student).selectinload(StudentProfile.branch)
    )

    if current_user.role == UserRole.STUDENT:
        stmt_prof = select(StudentProfile).where(StudentProfile.user_id == current_user.id)
        student = (await db.execute(stmt_prof)).scalars().first()
        if not student:
            return []
        stmt = stmt.where(AttendanceRecord.student_id == student.id)
    elif current_user.role == UserRole.TEACHER:
        stmt_t = select(TeacherProfile).where(TeacherProfile.user_id == current_user.id)
        teacher = (await db.execute(stmt_t)).scalars().first()
        if not teacher:
            return []
        stmt = stmt.join(AttendanceRecord.lecture_session).where(LectureSession.teacher_id == teacher.id)
    else:
        # Admin filters
        if course_id:
            stmt = stmt.join(AttendanceRecord.student).where(StudentProfile.course_id == course_id)
        if branch_id:
            stmt = stmt.join(AttendanceRecord.student).where(StudentProfile.branch_id == branch_id)

    if subject_id:
        stmt = stmt.join(AttendanceRecord.lecture_session).where(LectureSession.subject_id == subject_id)
    if status_filter:
        stmt = stmt.where(AttendanceRecord.status == status_filter)
    if date_from:
        stmt = stmt.where(AttendanceRecord.recorded_at >= datetime.combine(date_from, time(0, 0), tzinfo=timezone.utc))
    if date_to:
        stmt = stmt.where(AttendanceRecord.recorded_at <= datetime.combine(date_to, time(23, 59, 59), tzinfo=timezone.utc))

    stmt = stmt.order_by(AttendanceRecord.recorded_at.desc())
    return (await db.execute(stmt)).scalars().all()

@router.get("/excel")
async def export_attendance_excel(
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    subject_id: Optional[int] = None,
    status_filter: Optional[AttendanceStatus] = None,
    course_id: Optional[int] = None,
    branch_id: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    records = await fetch_filtered_records(db, current_user, date_from, date_to, subject_id, status_filter, course_id, branch_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    # Header styling
    headers = [
        "Record ID", "Student ID", "Student Name", "Roll Number",
        "Course", "Branch", "Subject Code", "Subject Name",
        "Date", "Recorded Time", "Status", "Method", "Remarks"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in records:
        stud = r.student
        sess = r.lecture_session
        ws.append([
            r.id,
            sanitize_formula(stud.user.user_code if stud and stud.user else ""),
            sanitize_formula(stud.user.full_name if stud and stud.user else ""),
            sanitize_formula(stud.roll_number if stud else ""),
            sanitize_formula(stud.course.name if stud and stud.course else ""),
            sanitize_formula(stud.branch.name if stud and stud.branch else ""),
            sanitize_formula(sess.subject.code if sess and sess.subject else ""),
            sanitize_formula(sess.subject.name if sess and sess.subject else ""),
            str(sess.date if sess else ""),
            r.recorded_at.strftime("%H:%M:%S") if r.recorded_at else "",
            r.status.value,
            r.verification_method.value,
            sanitize_formula(r.remarks or "")
        ])

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"attendx_report_{date.today().strftime('%Y%m%d')}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/csv")
async def export_attendance_csv(
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    subject_id: Optional[int] = None,
    status_filter: Optional[AttendanceStatus] = None,
    course_id: Optional[int] = None,
    branch_id: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    records = await fetch_filtered_records(db, current_user, date_from, date_to, subject_id, status_filter, course_id, branch_id)

    output = io.StringIO()
    # Write UTF-8 BOM for Excel
    output.write("\ufeff")
    writer = csv.writer(output)

    writer.writerow([
        "Record ID", "Student ID", "Student Name", "Roll Number",
        "Course", "Branch", "Subject Code", "Subject Name",
        "Date", "Recorded Time", "Status", "Method", "Remarks"
    ])

    for r in records:
        stud = r.student
        sess = r.lecture_session
        writer.writerow([
            r.id,
            sanitize_formula(stud.user.user_code if stud and stud.user else ""),
            sanitize_formula(stud.user.full_name if stud and stud.user else ""),
            sanitize_formula(stud.roll_number if stud else ""),
            sanitize_formula(stud.course.name if stud and stud.course else ""),
            sanitize_formula(stud.branch.name if stud and stud.branch else ""),
            sanitize_formula(sess.subject.code if sess and sess.subject else ""),
            sanitize_formula(sess.subject.name if sess and sess.subject else ""),
            str(sess.date if sess else ""),
            r.recorded_at.strftime("%H:%M:%S") if r.recorded_at else "",
            r.status.value,
            r.verification_method.value,
            sanitize_formula(r.remarks or "")
        ])

    filename = f"attendx_report_{date.today().strftime('%Y%m%d')}.csv"
    return Response(
        content=output.getvalue().encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/pdf")
async def export_attendance_pdf(
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    subject_id: Optional[int] = None,
    status_filter: Optional[AttendanceStatus] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    records = await fetch_filtered_records(db, current_user, date_from, date_to, subject_id, status_filter)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()

    story = []
    # Title
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=6
    )
    story.append(Paragraph("AttendX — Attendance Verification Report", title_style))
    story.append(Paragraph(f"Institution: GCRG Group of Institutions, Lucknow | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
    story.append(Spacer(1, 15))

    table_data = [
        ["ID", "Student ID", "Student Name", "Roll No", "Subject", "Date", "Time", "Status", "Method"]
    ]

    for r in records[:100]:  # Cap for PDF page limits
        stud = r.student
        sess = r.lecture_session
        table_data.append([
            str(r.id),
            stud.user.user_code if stud and stud.user else "",
            stud.user.full_name if stud and stud.user else "",
            stud.roll_number if stud else "",
            sess.subject.code if sess and sess.subject else "",
            str(sess.date if sess else ""),
            r.recorded_at.strftime("%H:%M") if r.recorded_at else "",
            r.status.value,
            r.verification_method.value
        ])

    if len(table_data) == 1:
        table_data.append(["-", "No attendance records available for selected criteria.", "", "", "", "", "", "", ""])

    t = Table(table_data)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F8FAFC')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('FONTSIZE', (0, 1), (-1, -1), 9)
    ]))

    story.append(t)
    doc.build(story)
    buffer.seek(0)

    filename = f"attendx_report_{date.today().strftime('%Y%m%d')}.pdf"
    return Response(
        content=buffer.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
